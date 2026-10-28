import openpyxl
from app.matching.inquiry import parse_inquiry_xlsx, InquiryLine

def _make_inquiry(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inquiry table"
    ws.append(["Remark", "Product", "Commodity", "Category", "Commodity\nManager",
               "Buyer", "Vendor list", "Descripton", "Qty1", "Qty.", "Unit\nPrice"])
    ws.append(["本次报价", "7-Axis七轴", "Equipment", "7-Axis", "Li", "Li",
               "推荐", "机器人七轴 含底座 11.5m", 2, 2, 47900])
    ws.append(["历史报价", "Robot 机器人", "Equipment", "Robot", "Li", "Li",
               "FANUC", "R-2000iC", 8, 8, 300000])
    f = tmp_path / "inquiry.xlsx"
    wb.save(f)
    return f

def test_parse_maps_headers_tolerantly(tmp_path):
    lines = parse_inquiry_xlsx(_make_inquiry(tmp_path))
    assert len(lines) == 2
    l0 = lines[0]
    assert isinstance(l0, InquiryLine)
    assert l0.row_index == 2
    assert "七轴" in l0.product and "11.5m" in l0.spec
    assert l0.qty == 2 and l0.brand is None
    assert lines[1].brand == "FANUC" and lines[1].category == "Robot"

def test_prefers_sheet_named_inquiry(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "汇总"
    ws.append(["杂项", "表头"]); ws.append([1, 2])
    ws2 = wb.create_sheet("Inquiry table")
    ws2.append(["Product", "Qty"]); ws2.append(["七轴", 1])
    f = tmp_path / "multi.xlsx"; wb.save(f)
    lines = parse_inquiry_xlsx(f)
    assert len(lines) == 1 and lines[0].product == "七轴"
