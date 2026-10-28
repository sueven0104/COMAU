import openpyxl
from app import db
from app.export.draft import export_draft
from app.matching.inquiry import InquiryLine
from app.matching.matcher import MatchResult

def test_export_draft_columns_and_rows(tmp_path):
    conn = db.connect(":memory:")
    fid = db.upsert_source_file(conn, {
        "path": "D:/BAIC/盛允.xlsx", "content_hash": "h", "file_type": ".xlsx",
        "project_code": "BAIC", "email_subject": None, "email_from": None,
        "email_date": None, "email_body": None, "processed_at": db.now()})
    ids = db.insert_quote_items(conn, fid, "BAIC", [
        {"supplier": "盛允", "category": "7-Axis", "product_name": "机器人七轴",
         "spec": "11.5m", "brand": "盛允", "model": None, "qty": 1, "unit": "套",
         "unit_price": 100.0, "currency": "CNY", "tax_included": True,
         "price_status": "成交价", "quote_date": "2026-09-01", "terms": None,
         "source_location": "R2", "confidence": 0.9, "needs_review": 0}])
    lines = [InquiryLine(2, "7-Axis七轴", "机器人七轴 11.5m", "盛允", 2)]
    matches = [MatchResult(2, "精确匹配", suggested_price=100.0,
                           quote_item_ids=ids, basis="BAIC/盛允/2026-09-01/成交价",
                           explanation="参考成交价")]
    out = tmp_path / "draft.xlsx"
    export_draft(conn, lines, matches, out)
    ws = openpyxl.load_workbook(out).active
    header = [c.value for c in ws[1]]
    assert header[:3] == ["行号", "产品", "规格"]
    assert "建议价" in header and "溯源" in header and "需人工确认" in header
    row2 = [c.value for c in ws[2]]
    assert row2[0] == 2 and row2[header.index("建议价")] == 100.0
    assert "盛允.xlsx" in row2[header.index("溯源")]
