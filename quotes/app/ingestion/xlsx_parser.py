"""Excel 全 sheet 序列化为文本（结构判断交给 LLM 抽取）。"""
import datetime
from dataclasses import dataclass
from pathlib import Path
import openpyxl

@dataclass
class SheetData:
    name: str
    rows: list[list[str]]

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):  # datetime 是 date 的子类，须先判
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)

def parse_xlsx(path: Path, max_rows: int = 400, max_cols: int = 40) -> list[SheetData]:
    """逐 sheet 读取为字符串二维表。超过 max_rows/max_cols 的行列静默截断；
    全空 sheet 丢弃。无法加载的损坏文件（BadZipFile 等）向上传播，由管道按文件捕获。"""
    sheets: list[SheetData] = []
    # openpyxl 3.1 的 Workbook 不支持上下文管理器协议，用 try/finally 保证
    # 迭代中途抛异常时也释放句柄（Windows 上否则锁住文件）
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for row in ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
                cells = [_cell(c) for c in row[:max_cols]]
                # read_only 模式下 iter_rows 会 pad 到 max_cols，裁掉行尾空单元格
                last = max((i for i, c in enumerate(cells) if c.strip()), default=-1)
                cells = cells[:last + 1]
                if cells:
                    rows.append(cells)
            if rows:
                sheets.append(SheetData(ws.title, rows))
    finally:
        wb.close()
    return sheets
