"""报价建议草稿导出：原清单行 + 建议价 + 档位 + 依据 + 溯源 + 人工确认标记。"""
import logging
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

log = logging.getLogger(__name__)

COLUMNS = ["行号", "产品", "规格", "品牌", "数量", "建议价", "价格区间",
           "匹配档位", "参考依据", "差异说明", "溯源", "需人工确认"]

NEED_CONFIRM = {"弱匹配", "无匹配"}

def export_draft(conn, lines, matches, out_path):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价建议草稿"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EFEFF2")
    for ln, m in zip(lines, matches):
        trace = ""
        if m.quote_item_ids:
            rows = conn.execute(
                """SELECT sf.path, qi.source_location FROM quote_items qi
                   JOIN source_files sf ON sf.id=qi.source_file_id WHERE qi.id=?""",
                (m.quote_item_ids[0],)).fetchall()
            if rows:
                p = Path(rows[0]["path"])
                trace = f"{p.name} {rows[0]['source_location'] or ''}".strip()
        ws.append([
            ln.row_index, ln.product, ln.spec, ln.brand, ln.qty,
            m.suggested_price,
            f"{m.price_range[0]:.0f}~{m.price_range[1]:.0f}" if m.price_range else None,
            m.tier, m.basis, m.explanation, trace,
            "是" if m.tier in NEED_CONFIRM else ""])
    for col, width in zip("ABCDEFGHIJKL", [6, 22, 28, 12, 8, 12, 16, 10, 40, 36, 30, 10]):
        ws.column_dimensions[col].width = width
    wb.save(out_path)
    n_confirm = sum(1 for m in matches if m.tier in NEED_CONFIRM)
    log.info("草稿导出 %s：共 %d 行，其中需人工确认 %d 行", out_path, len(lines), n_confirm)
    return out_path
